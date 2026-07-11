"""
LeadScraper Pro v2.0 – Complete Flask Backend
All-in-one: scraping engine, REST API, WebSocket, CSV/Excel export.
"""

import os
import re
import csv
import io
import html
import uuid
import threading
from datetime import datetime

import requests
from bs4 import BeautifulSoup
import phonenumbers
from flask import Flask, render_template, request, jsonify, send_file
from flask_socketio import SocketIO, emit
from flask_cors import CORS
import openpyxl

# ---------------------------------------------------------------------------
# App Setup
# ---------------------------------------------------------------------------
app = Flask(__name__, template_folder=".", static_folder=".")
CORS(app)
socketio = SocketIO(app, cors_allowed_origins="*", async_mode="threading")

# ---------------------------------------------------------------------------
# Thread-safe in-memory stores
# ---------------------------------------------------------------------------
_lock = threading.Lock()
jobs: dict = {}
leads_db: dict = {}

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36",
    "Accept-Language": "en-US,en;q=0.9",
}
EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"[\+]?[\d\s\-\(\)]{7,20}")


def _normalise_phone(raw: str, region: str = "US") -> str:
    try:
        p = phonenumbers.parse(raw, region)
        if phonenumbers.is_valid_number(p):
            return phonenumbers.format_number(p, phonenumbers.PhoneNumberFormat.E164)
    except Exception:
        pass
    return raw.strip()


def _extract_emails(text: str) -> list[str]:
    return list({m.lower() for m in EMAIL_RE.findall(text)})


def _extract_phones(text: str) -> list[str]:
    seen, out = set(), []
    for r in PHONE_RE.findall(text):
        n = _normalise_phone(r)
        if n and n not in seen:
            seen.add(n)
            out.append(n)
    return out


def _fetch(url: str, timeout: int = 15) -> requests.Response | None:
    try:
        r = requests.get(url, headers=HEADERS, timeout=timeout, allow_redirects=True)
        return r if r.ok else None
    except Exception:
        return None


def _sanitize(text: str) -> str:
    """Escape HTML entities to prevent XSS."""
    return html.escape(text).strip()


# ---------------------------------------------------------------------------
# Scrapers
# ---------------------------------------------------------------------------
def scrape_website_contacts(url: str) -> dict:
    resp = _fetch(url)
    if not resp:
        return {"url": url, "emails": [], "phones": [], "socials": [], "error": "fetch_failed"}
    soup = BeautifulSoup(resp.text, "lxml")
    text = soup.get_text(" ", strip=True)
    socials = []
    social_domains = ["facebook.com", "twitter.com", "x.com", "linkedin.com",
                      "instagram.com", "youtube.com", "tiktok.com"]
    for a in soup.find_all("a", href=True):
        if any(d in a["href"] for d in social_domains):
            socials.append(a["href"])
    return {"url": url, "emails": _extract_emails(text), "phones": _extract_phones(text),
            "socials": list(set(socials))}


def scrape_google_maps(query: str, max_results: int = 20) -> list[dict]:
    results = []
    resp = _fetch(f"https://www.google.com/search?q={requests.utils.quote(query)}&tbm=lcl")
    if not resp:
        return results
    soup = BeautifulSoup(resp.text, "lxml")
    for item in soup.select("div[data-attrid]"):
        name_el = item.select_one("div[role='heading']")
        if not name_el:
            continue
        name = name_el.get_text(strip=True)
        addr_el = item.select_one("div[data-attrid*='address']")
        phone_el = item.select_one("div[data-attrid*='phone']")
        web_el = item.select_one("a[data-attrid*='website']")
        rating_el = item.select_one("span[aria-label]")
        if name:
            results.append({
                "id": str(uuid.uuid4())[:8],
                "name": _sanitize(name),
                "address": _sanitize(addr_el.get_text(strip=True)) if addr_el else "",
                "phone": _normalise_phone(phone_el.get_text(strip=True)) if phone_el else "",
                "website": web_el["href"] if web_el else "",
                "rating": rating_el["aria-label"] if rating_el else "",
                "source": "google_maps",
                "scraped_at": datetime.utcnow().isoformat(),
            })
        if len(results) >= max_results:
            break
    return results


def scrape_website_bulk(query: str, max_results: int = 20) -> list[dict]:
    """If query is a URL, scrape it for contacts."""
    results = []
    if query.startswith("http"):
        r = scrape_website_contacts(query)
        if not r.get("error"):
            results.append({
                "id": str(uuid.uuid4())[:8],
                "name": _sanitize(r["url"]),
                "emails": r["emails"],
                "phones": r["phones"],
                "socials": r["socials"],
                "source": "website_contacts",
                "scraped_at": datetime.utcnow().isoformat(),
            })
    return results


def scrape_linkedin_people(query: str, max_results: int = 20) -> list[dict]:
    results = []
    resp = _fetch(f"https://www.linkedin.com/search/results/people/?keywords={requests.utils.quote(query)}")
    if not resp:
        return results
    soup = BeautifulSoup(resp.text, "lxml")
    for card in soup.select("li.reusable-search__result-container"):
        name_el = card.select_one("span.entity-result__title-text a")
        if not name_el:
            continue
        name = name_el.get_text(strip=True)
        if name:
            headline_el = card.select_one("div.entity-result__primary-subtitle")
            location_el = card.select_one("div.entity-result__secondary-subtitle")
            results.append({
                "id": str(uuid.uuid4())[:8],
                "name": _sanitize(name),
                "headline": _sanitize(headline_el.get_text(strip=True)) if headline_el else "",
                "location": _sanitize(location_el.get_text(strip=True)) if location_el else "",
                "profile_url": name_el.get("href", ""),
                "source": "linkedin",
                "scraped_at": datetime.utcnow().isoformat(),
            })
        if len(results) >= max_results:
            break
    return results


# ---------------------------------------------------------------------------
# Background Job Runner (thread-safe)
# ---------------------------------------------------------------------------
SCRAPERS = {
    "google_maps": scrape_google_maps,
    "website_contacts": scrape_website_bulk,
    "linkedin": scrape_linkedin_people,
}


def _run_job(job_id: str):
    job = jobs[job_id]
    job["status"] = "running"
    job["started_at"] = datetime.utcnow().isoformat()
    socketio.emit("job_update", {"job_id": job_id, "status": "running"})

    all_leads: list[dict] = []
    sources = job.get("sources", [])
    query = job.get("query", "")
    max_results = job.get("max_results", 20)

    for i, source in enumerate(sources):
        with _lock:
            job["current_source"] = source
            job["progress"] = int((i / len(sources)) * 100)
        socketio.emit("job_update", {"job_id": job_id, "progress": job["progress"], "current_source": source})

        scraper = SCRAPERS.get(source)
        if scraper:
            try:
                leads = scraper(query, max_results)
                all_leads.extend(leads)
            except Exception as e:
                socketio.emit("job_error", {"job_id": job_id, "error": f"{source}: {e}", "source": source})

    with _lock:
        leads_db[job_id] = all_leads
        job["status"] = "completed"
        job["progress"] = 100
        job["completed_at"] = datetime.utcnow().isoformat()
        job["total_leads"] = len(all_leads)

    socketio.emit("job_update", {"job_id": job_id, "status": "completed", "progress": 100, "total_leads": len(all_leads)})


# ---------------------------------------------------------------------------
# Routes – Frontend
# ---------------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


# ---------------------------------------------------------------------------
# Routes – API
# ---------------------------------------------------------------------------
@app.route("/api/jobs", methods=["POST"])
def create_job():
    data = request.get_json(force=True)
    query = _sanitize(data.get("query", "").strip())
    sources = data.get("sources", ["google_maps"])
    max_results = min(int(data.get("max_results", 20)), 100)
    if not query:
        return jsonify({"error": "query is required"}), 400
    job_id = str(uuid.uuid4())[:12]
    with _lock:
        jobs[job_id] = {"id": job_id, "query": query, "sources": sources, "max_results": max_results,
                         "status": "pending", "progress": 0, "created_at": datetime.utcnow().isoformat()}
    threading.Thread(target=_run_job, args=(job_id,), daemon=True).start()
    return jsonify({"job_id": job_id, "status": "pending"}), 201


@app.route("/api/jobs", methods=["GET"])
def list_jobs():
    with _lock:
        return jsonify(list(jobs.values()))


@app.route("/api/jobs/<job_id>", methods=["GET"])
def get_job(job_id: str):
    with _lock:
        job = jobs.get(job_id)
    if not job:
        return jsonify({"error": "not found"}), 404
    data = dict(job)
    with _lock:
        data["leads"] = leads_db.get(job_id, [])
    return jsonify(data)


@app.route("/api/jobs/<job_id>/leads", methods=["GET"])
def get_leads(job_id: str):
    with _lock:
        leads = leads_db.get(job_id, [])
    return jsonify({"job_id": job_id, "count": len(leads), "leads": leads})


@app.route("/api/jobs/<job_id>/export/csv", methods=["GET"])
def export_csv(job_id: str):
    with _lock:
        leads = leads_db.get(job_id, [])
    if not leads:
        return jsonify({"error": "no leads found"}), 404
    all_keys = list(dict.fromkeys(k for lead in leads for k in lead))
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=all_keys)
    w.writeheader()
    for lead in leads:
        row = {k: "; ".join(str(i) for i in v) if isinstance(v, list) else v for k, v in lead.items()}
        w.writerow(row)
    buf.seek(0)
    return send_file(io.BytesIO(buf.getvalue().encode("utf-8")), mimetype="text/csv",
                     as_attachment=True, download_name=f"leads_{job_id}.csv")


@app.route("/api/jobs/<job_id>/export/xlsx", methods=["GET"])
def export_xlsx(job_id: str):
    with _lock:
        leads = leads_db.get(job_id, [])
    if not leads:
        return jsonify({"error": "no leads found"}), 404
    all_keys = list(dict.fromkeys(k for lead in leads for k in lead))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"
    for col, key in enumerate(all_keys, 1):
        c = ws.cell(row=1, column=col, value=key.replace("_", " ").title())
        c.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        c.fill = openpyxl.styles.PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for ri, lead in enumerate(leads, 2):
        for ci, key in enumerate(all_keys, 1):
            val = lead.get(key, "")
            ws.cell(row=ri, column=ci, value="; ".join(str(i) for i in val) if isinstance(val, list) else val)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max((len(str(c.value or "")) for c in col), default=10) + 2, 40)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=f"leads_{job_id}.xlsx")


@app.route("/api/stats")
def get_stats():
    with _lock:
        total_jobs = len(jobs)
        total_leads = sum(len(v) for v in leads_db.values())
        running = sum(1 for j in jobs.values() if j["status"] == "running")
        completed = sum(1 for j in jobs.values() if j["status"] == "completed")
    return jsonify({"total_jobs": total_jobs, "total_leads": total_leads,
                    "running_jobs": running, "completed_jobs": completed})


# ---------------------------------------------------------------------------
# WebSocket
# ---------------------------------------------------------------------------
@socketio.on("connect")
def handle_connect():
    emit("connected", {"status": "ok"})


@socketio.on("subscribe_job")
def handle_subscribe(data):
    job_id = data.get("job_id")
    with _lock:
        job = jobs.get(job_id) if job_id else None
    if job:
        emit("subscribed", {"job_id": job_id, "status": job["status"]})


# ---------------------------------------------------------------------------
# Run
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    print("\n🚀 LeadScraper Pro v2.0 starting on http://localhost:5000\n")
    socketio.run(app, host="0.0.0.0", port=5000, debug=True)
